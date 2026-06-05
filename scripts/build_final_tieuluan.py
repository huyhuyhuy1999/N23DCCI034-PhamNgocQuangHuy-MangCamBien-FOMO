from __future__ import annotations

import math
import os
import textwrap
from pathlib import Path
from typing import Iterable, Sequence
from urllib.request import urlretrieve

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


SCRIPT_DIR = Path(__file__).resolve().parent
if (SCRIPT_DIR / "edge_impulse_screenshots").exists() or (SCRIPT_DIR / "HuongDan_NopTieuLuan_CuoiKy.docx").exists():
    ROOT = SCRIPT_DIR
elif (SCRIPT_DIR.parent / "screenshots").exists() or (SCRIPT_DIR.parent / "report").exists():
    ROOT = SCRIPT_DIR.parent
else:
    ROOT = Path(os.environ.get("MCB_ROOT", str(SCRIPT_DIR)))

DEFAULT_OUT_DOCX = ROOT / "PhamNgocQuangHuy_final_cuoiky.docx"
if (ROOT / "report").exists():
    DEFAULT_OUT_DOCX = ROOT / "report" / "PhamNgocQuangHuy_final_cuoiky.docx"
OUT_DOCX = Path(os.environ.get("OUT_DOCX", str(DEFAULT_OUT_DOCX)))
ASSET_DIR = ROOT / "_generated_tieuluan_assets"
ASSET_DIR.mkdir(exist_ok=True)
EI_SCREEN_DIR = ROOT / "edge_impulse_screenshots"
if not EI_SCREEN_DIR.exists() and (ROOT / "screenshots").exists():
    EI_SCREEN_DIR = ROOT / "screenshots"

PYTHON_DEPS = Path(r"C:\Users\NK\.cache\codex-runtimes\codex-primary-runtime\dependencies\python")
FONT_CANDIDATES = [
    Path(r"C:\Windows\Fonts\times.ttf"),
    Path(r"C:\Windows\Fonts\timesbd.ttf"),
    PYTHON_DEPS / "matplotlib" / "mpl-data" / "fonts" / "ttf" / "DejaVuSans.ttf",
]

PTIT_LOGO_URL = "https://commons.wikimedia.org/wiki/Special:Redirect/file/Logo_PTIT_University.png"
PTIT_LOGO = ASSET_DIR / "ptit_logo.png"


def find_font(bold: bool = False) -> str | None:
    preferred = Path(r"C:\Windows\Fonts\timesbd.ttf" if bold else r"C:\Windows\Fonts\times.ttf")
    if preferred.exists():
        return str(preferred)
    for candidate in FONT_CANDIDATES:
        if candidate.exists():
            return str(candidate)
    return None


def pil_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    path = find_font(bold)
    if path:
        return ImageFont.truetype(path, size=size)
    return ImageFont.load_default()


def set_run_font(run, size: int | float | None = None, bold: bool | None = None,
                 italic: bool | None = None, color: str | None = None,
                 name: str = "Times New Roman") -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_style_font(style, size: int | float, bold: bool = False, color: str | None = None,
                   name: str = "Times New Roman") -> None:
    style.font.name = name
    style._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    style.font.size = Pt(size)
    style.font.bold = bold
    if color:
        style.font.color.rgb = RGBColor.from_string(color)


def configure_section(section, left=3.0, right=2.0, top=2.0, bottom=2.0) -> None:
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(left)
    section.right_margin = Cm(right)
    section.top_margin = Cm(top)
    section.bottom_margin = Cm(bottom)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    set_style_font(normal, 13)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Cm(1.27)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_before = Pt(6)
    normal.paragraph_format.space_after = Pt(0)

    for style_name, size, align, before, after in [
        ("Heading 1", 16, WD_ALIGN_PARAGRAPH.CENTER, 6, 18),
        ("Heading 2", 14, WD_ALIGN_PARAGRAPH.LEFT, 6, 6),
        ("Heading 3", 13, WD_ALIGN_PARAGRAPH.LEFT, 6, 6),
    ]:
        style = styles[style_name]
        set_style_font(style, size, True)
        style.paragraph_format.alignment = align
        style.paragraph_format.line_spacing = 1.5 if style_name == "Heading 1" else 1.2
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    for style_name in ["List Bullet", "List Number"]:
        style = styles[style_name]
        set_style_font(style, 13)
        style.paragraph_format.left_indent = Cm(1.27)
        style.paragraph_format.first_line_indent = Cm(-0.5)
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.space_before = Pt(3)
        style.paragraph_format.space_after = Pt(0)


def add_field(paragraph, instr: str) -> None:
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = instr
    fld_char_sep = OxmlElement("w:fldChar")
    fld_char_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_sep)
    run._r.append(text)
    run._r.append(fld_char_end)


def set_page_number_format(section, fmt: str = "decimal", start: int = 1) -> None:
    sect_pr = section._sectPr
    pg_num_type = sect_pr.find(qn("w:pgNumType"))
    if pg_num_type is None:
        pg_num_type = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num_type)
    pg_num_type.set(qn("w:fmt"), fmt)
    pg_num_type.set(qn("w:start"), str(start))


def add_footer_page_number(section, fmt: str = "decimal", start: int = 1, prefix: str = "") -> None:
    set_page_number_format(section, fmt=fmt, start=start)
    section.footer.is_linked_to_previous = False
    for p in section.footer.paragraphs:
        p.text = ""
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if prefix:
        r = p.add_run(prefix)
        set_run_font(r, 11)
    add_field(p, "PAGE")


def clear_footer(section) -> None:
    section.footer.is_linked_to_previous = False
    for p in section.footer.paragraphs:
        p.text = ""


def set_table_width(table, widths_cm: Sequence[float]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in table.rows:
        for idx, width in enumerate(widths_cm):
            cell = row.cells[idx]
            cell.width = Cm(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(Cm(width).twips)))
            tc_w.set(qn("w:type"), "dxa")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def format_table(table, widths_cm: Sequence[float] | None = None, header_rows: int = 1) -> None:
    if widths_cm:
        set_table_width(table, widths_cm)
    for r_idx, row in enumerate(table.rows):
        if r_idx < header_rows:
            repeat_table_header(row)
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            for p in cell.paragraphs:
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    set_run_font(run, 11.5)
            if r_idx < header_rows:
                set_cell_shading(cell, "E8EEF5")
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.bold = True
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_para(doc: Document, text: str = "", style: str | None = None,
             align: int | None = None, bold: bool = False,
             italic: bool = False, first_line: bool = True,
             size: int | float | None = None) -> None:
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if not first_line:
        p.paragraph_format.first_line_indent = Cm(0)
    if text:
        r = p.add_run(text)
        set_run_font(r, size=size or (13 if style is None else None), bold=bold, italic=italic)


def add_multirun_para(doc: Document, runs: Sequence[tuple[str, bool, bool]], style: str | None = None,
                      align=None, first_line: bool = True) -> None:
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if not first_line:
        p.paragraph_format.first_line_indent = Cm(0)
    for text, bold, italic in runs:
        r = p.add_run(text)
        set_run_font(r, 13, bold=bold, italic=italic)


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(text)
    set_run_font(r, 13)


def add_number(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Number")
    r = p.add_run(text)
    set_run_font(r, 13)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    p.paragraph_format.first_line_indent = Cm(0)
    for run in p.runs:
        set_run_font(run, {1: 16, 2: 14, 3: 13}.get(level, 13), bold=True)


def add_page_break(doc: Document) -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_caption(doc: Document, text: str, kind: str = "figure") -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(text)
    set_run_font(r, 12, italic=(kind == "figure"))


def add_table_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_run_font(r, 12, bold=True)


def add_placeholder(doc: Document, title: str, detail: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    format_table(table, [16.0], header_rows=0)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "FFF7E6")
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    r = p.add_run(title)
    set_run_font(r, 12, bold=True, color="7A4D00")
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.first_line_indent = Cm(0)
    r2 = p2.add_run(detail)
    set_run_font(r2, 11.5, italic=True, color="7A4D00")


def add_screenshot(doc: Document, file_name: str, caption: str, width_cm: float = 15.5) -> bool:
    path = EI_SCREEN_DIR / file_name
    if not path.exists():
        return False
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    add_caption(doc, caption)
    return True


def draw_rounded_rect(draw: ImageDraw.ImageDraw, xy, radius, fill, outline=None, width=1):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)


def center_text(draw: ImageDraw.ImageDraw, box, text: str, font, fill=(20, 20, 20), line_gap=8) -> None:
    x1, y1, x2, y2 = box
    lines = []
    for raw in text.split("\n"):
        if raw:
            lines.extend(textwrap.wrap(raw, width=24))
        else:
            lines.append("")
    heights = []
    widths = []
    for line in lines:
        bb = draw.textbbox((0, 0), line, font=font)
        widths.append(bb[2] - bb[0])
        heights.append(bb[3] - bb[1])
    total_h = sum(heights) + line_gap * (len(lines) - 1)
    y = y1 + (y2 - y1 - total_h) / 2
    for line, w, h in zip(lines, widths, heights):
        draw.text((x1 + (x2 - x1 - w) / 2, y), line, font=font, fill=fill)
        y += h + line_gap


def arrow(draw: ImageDraw.ImageDraw, start, end, fill=(75, 85, 99), width=5) -> None:
    draw.line([start, end], fill=fill, width=width)
    sx, sy = start
    ex, ey = end
    angle = math.atan2(ey - sy, ex - sx)
    size = 16
    pts = [
        (ex, ey),
        (ex - size * math.cos(angle - math.pi / 6), ey - size * math.sin(angle - math.pi / 6)),
        (ex - size * math.cos(angle + math.pi / 6), ey - size * math.sin(angle + math.pi / 6)),
    ]
    draw.polygon(pts, fill=fill)


def make_diagram_architecture() -> Path:
    path = ASSET_DIR / "diagram_architecture.png"
    img = Image.new("RGB", (1600, 850), "white")
    d = ImageDraw.Draw(img)
    title_font = pil_font(42, True)
    label_font = pil_font(30, True)
    small_font = pil_font(24)
    d.text((60, 35), "Kiến trúc hệ thống phát hiện chai nhựa tái chế", font=title_font, fill=(11, 37, 69))
    boxes = [
        ((70, 190, 350, 390), "Camera\nkhu vực phân loại", "#E3F2FD"),
        ((440, 190, 720, 390), "Edge node\nESP32-CAM / Raspberry Pi", "#E8F5E9"),
        ((810, 190, 1090, 390), "FOMO model\n500ml - 1L - 1.5L", "#FFF3E0"),
        ((1180, 190, 1460, 390), "Dashboard / log\nđếm và cảnh báo", "#F3E5F5"),
    ]
    for box, label, fill in boxes:
        draw_rounded_rect(d, box, 28, fill, (91, 106, 125), 3)
        center_text(d, box, label, label_font)
    for x in [350, 720, 1090]:
        arrow(d, (x + 20, 290), (x + 80, 290))
    lower = [
        ((175, 540, 470, 710), "Ảnh đầu vào\n96x96/160x160", "#F8FAFC"),
        ((620, 540, 980, 710), "Tiền xử lý\nresize, crop, RGB/grayscale", "#F8FAFC"),
        ((1120, 540, 1420, 710), "Kết quả\nnhãn + centroid + độ tin cậy", "#F8FAFC"),
    ]
    for box, label, fill in lower:
        draw_rounded_rect(d, box, 24, fill, (148, 163, 184), 2)
        center_text(d, box, label, small_font)
    arrow(d, (470, 625), (620, 625), width=4)
    arrow(d, (980, 625), (1120, 625), width=4)
    img.save(path)
    return path


def make_diagram_pipeline() -> Path:
    path = ASSET_DIR / "diagram_pipeline.png"
    img = Image.new("RGB", (1600, 760), "white")
    d = ImageDraw.Draw(img)
    title_font = pil_font(40, True)
    label_font = pil_font(25, True)
    small_font = pil_font(21)
    d.text((60, 35), "Pipeline Edge Impulse cho đề tài", font=title_font, fill=(11, 37, 69))
    labels = [
        "Project Edge Impulse\n500ml - 1L - 1.5L",
        "Thu ảnh chai\n500ml, 1L, 1.5L",
        "Gán bounding box\nvà nhãn size",
        "Create impulse\nImages + Object Detection",
        "Huấn luyện FOMO\nMobileNetV2",
        "Deploy\nWebAssembly / C++",
    ]
    x = 65
    for idx, label in enumerate(labels):
        box = (x, 210, x + 210, 415)
        draw_rounded_rect(d, box, 26, "#EEF6FF" if idx % 2 == 0 else "#F0FDF4", (71, 85, 105), 2)
        center_text(d, box, label, label_font)
        if idx < len(labels) - 1:
            arrow(d, (x + 220, 312), (x + 300, 312), width=4)
        x += 255
    d.text((95, 535), "Điểm kiểm tra bắt buộc: dataset cân bằng, nhãn nhất quán, confusion matrix, F1 score, kiểm thử realtime trên trình duyệt.", font=small_font, fill=(51, 65, 85))
    d.text((95, 590), "Các ảnh minh chứng thật sẽ được thay vào các khung placeholder trong báo cáo trước khi nộp.", font=small_font, fill=(127, 76, 0))
    img.save(path)
    return path


def make_diagram_inference() -> Path:
    path = ASSET_DIR / "diagram_inference.png"
    img = Image.new("RGB", (1600, 820), "white")
    d = ImageDraw.Draw(img)
    title_font = pil_font(40, True)
    label_font = pil_font(26, True)
    small_font = pil_font(22)
    d.text((60, 35), "Luồng suy luận FOMO và ý nghĩa centroid", font=title_font, fill=(11, 37, 69))
    # Left pseudo-image
    frame = (95, 165, 620, 650)
    draw_rounded_rect(d, frame, 18, "#F8FAFC", (100, 116, 139), 3)
    d.text((135, 185), "Frame camera", font=label_font, fill=(30, 41, 59))
    bottles = [
        (180, 285, 245, 560, "500ml", "#A7F3D0"),
        (335, 245, 420, 565, "1L", "#BFDBFE"),
        (500, 215, 595, 575, "1.5L", "#FDE68A"),
    ]
    for x1, y1, x2, y2, label, fill in bottles:
        draw_rounded_rect(d, (x1, y1, x2, y2), 22, fill, (51, 65, 85), 3)
        cx, cy = (x1 + x2) // 2, (y1 + y2) // 2
        d.ellipse((cx - 8, cy - 8, cx + 8, cy + 8), fill=(220, 38, 38))
        d.text((x1 + 5, y2 + 12), label, font=small_font, fill=(30, 41, 59))
    arrow(d, (660, 400), (805, 400), width=5)
    middle = (835, 230, 1155, 570)
    draw_rounded_rect(d, middle, 24, "#FFF7ED", (194, 65, 12), 3)
    center_text(d, middle, "FOMO\nheatmap theo lưới\nphân lớp tại từng ô", label_font)
    arrow(d, (1185, 400), (1320, 400), width=5)
    result = (1345, 205, 1540, 595)
    draw_rounded_rect(d, result, 24, "#ECFDF5", (22, 101, 52), 3)
    center_text(d, result, "Output\nlabel\ncentroid\nscore", label_font)
    d.text((110, 705), "Bounding box dùng trong lúc gán nhãn và phân tích kích thước; khi triển khai FOMO, đầu ra chính là tâm đối tượng (centroid), không phải bbox đầy đủ.", font=small_font, fill=(127, 76, 0))
    img.save(path)
    return path


def make_diagram_metric() -> Path:
    path = ASSET_DIR / "diagram_metric.png"
    img = Image.new("RGB", (1500, 760), "white")
    d = ImageDraw.Draw(img)
    title_font = pil_font(38, True)
    label_font = pil_font(24, True)
    small_font = pil_font(21)
    d.text((55, 35), "Bounding-box size metric cho dữ liệu huấn luyện", font=title_font, fill=(11, 37, 69))
    canvas = (80, 155, 760, 610)
    draw_rounded_rect(d, canvas, 18, "#F8FAFC", (100, 116, 139), 3)
    d.text((105, 180), "Ảnh gốc: image_width x image_height", font=small_font, fill=(51, 65, 85))
    bbox = (250, 260, 570, 530)
    draw_rounded_rect(d, bbox, 10, "#DBEAFE", (37, 99, 235), 5)
    d.text((320, 375), "bbox", font=label_font, fill=(30, 64, 175))
    d.line([(250, 545), (570, 545)], fill=(37, 99, 235), width=4)
    d.text((335, 555), "bbox_width", font=small_font, fill=(30, 64, 175))
    d.line([(590, 260), (590, 530)], fill=(37, 99, 235), width=4)
    d.text((605, 380), "bbox_height", font=small_font, fill=(30, 64, 175))
    formula_box = (835, 205, 1410, 545)
    draw_rounded_rect(d, formula_box, 24, "#FFF7ED", (194, 65, 12), 3)
    center_text(
        d,
        formula_box,
        "bbox_area_ratio =\n(bbox_width x bbox_height) /\n(image_width x image_height)\n\nSo sánh mean/median\ncho 500ml, 1L, 1.5L",
        label_font,
    )
    img.save(path)
    return path


def ensure_logo() -> Path | None:
    if PTIT_LOGO.exists():
        return PTIT_LOGO
    try:
        urlretrieve(PTIT_LOGO_URL, PTIT_LOGO)
        return PTIT_LOGO
    except Exception:
        return None


def cover_page(doc: Document, secondary: bool = False) -> None:
    logo = ensure_logo()
    for text, size, bold in [
        ("BỘ KHOA HỌC VÀ CÔNG NGHỆ", 13, True),
        ("HỌC VIỆN CÔNG NGHỆ BƯU CHÍNH VIỄN THÔNG", 13, True),
        ("CƠ SỞ TẠI TP. HỒ CHÍ MINH", 13, True),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        r = p.add_run(text)
        set_run_font(r, size, bold=bold)
    add_para(doc, "", first_line=False)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    if logo:
        p.add_run().add_picture(str(logo), width=Cm(3.2))
    else:
        r = p.add_run("[LOGO PTIT]")
        set_run_font(r, 14, bold=True, color="9B1C1C")
    add_para(doc, "", first_line=False)
    for text, size in [
        ("TIỂU LUẬN MẠNG CẢM BIẾN", 18),
        ("NGÀNH CÔNG NGHỆ THÔNG TIN", 16),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        r = p.add_run(text)
        set_run_font(r, size, bold=True)
    add_para(doc, "", first_line=False)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    r = p.add_run("PHÁT HIỆN CHAI NHỰA TÁI CHẾ THEO SIZE\n(500ml, 1L, 1.5L) - NHÃN DÁN: 500ml, 1L, 1.5L")
    set_run_font(r, 16, bold=True)
    add_para(doc, "", first_line=False)
    info = [
        ("GVHD:", "Hồ Nhựt Minh"),
        ("SVTH:" if not secondary else "TÊN SINH VIÊN:", "Phạm Ngọc Quang Huy"),
        ("MSSV:", "N23DCCI034"),
        ("LỚP:", "D23CQCI01-N"),
        ("MÔN:", "Mạng Cảm Biến"),
    ]
    table = doc.add_table(rows=len(info), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_width(table, [3.5, 8.5])
    for i, (label, value) in enumerate(info):
        c0, c1 = table.rows[i].cells
        c0.text = label
        c1.text = value
        for cell in (c0, c1):
            for p in cell.paragraphs:
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.line_spacing = 1.2
                for run in p.runs:
                    set_run_font(run, 14, bold=(cell is c0))
    add_para(doc, "", first_line=False)
    add_para(doc, "", first_line=False)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    r = p.add_run("THÀNH PHỐ HỒ CHÍ MINH, THÁNG 06 NĂM 2026")
    set_run_font(r, 13, bold=True)


def task_page(doc: Document) -> None:
    add_heading(doc, "NHIỆM VỤ TIỂU LUẬN MẠNG CẢM BIẾN", 1)
    add_para(doc, "1. Họ và tên sinh viên được giao đề tài: Phạm Ngọc Quang Huy", first_line=False)
    add_para(doc, "MSSV: N23DCCI034        Lớp: D23CQCI01-N", first_line=False)
    add_para(doc, "2. Tên đề tài:", first_line=False)
    add_para(doc, "Phát hiện chai nhựa tái chế theo size (500ml, 1L, 1.5L) - Nhãn dán: 500ml, 1L, 1.5L.", first_line=False)
    add_para(doc, "3. Nhiệm vụ của đề tài:", first_line=False)
    for item in [
        "Tìm hiểu Edge AI, TinyML, mạng cảm biến có camera và mô hình object detection FOMO trên Edge Impulse.",
        "Tạo project Edge Impulse cho bài toán nhận diện chai nhựa theo ba nhãn 500ml, 1L và 1.5L.",
        "Xây dựng bộ nhãn 500ml, 1L, 1.5L; tạo dataset prototype có bounding box nhất quán để chạy pipeline và đề xuất hướng thu thập ảnh thật ở nhiều góc chụp, nền, ánh sáng và khoảng cách.",
        "Gán nhãn bằng bounding box, tính metric kích thước bbox_area_ratio để kiểm tra tính tách biệt kích thước giữa các class.",
        "Thiết kế impulse gồm xử lý ảnh và learning block Object Detection (FOMO), huấn luyện và đánh giá bằng confusion matrix/F1 score.",
        "Triển khai thử nghiệm WebAssembly trên trình duyệt và mô tả hướng tích hợp lên edge node như ESP32-CAM hoặc Raspberry Pi.",
        "Chèn screenshot minh chứng thật cho dataset, feature generation, huấn luyện, model testing, deployment build và browser realtime; ghi rõ realtime hiện mới chứng minh runtime webcam, chưa kiểm thử với chai thật ngoài hiện trường.",
    ]:
        add_bullet(doc, item)
    add_para(doc, "4. Link project Edge Impulse: https://studio.edgeimpulse.com/studio/1021052", first_line=False)
    add_para(doc, "5. Link repository bài nộp: https://github.com/<username>/<repository> (thay bằng link thật trước khi nộp).", first_line=False)


def declaration_pages(doc: Document) -> None:
    add_heading(doc, "LỜI CAM ĐOAN", 1)
    paragraphs = [
        "Tôi xin cam đoan đây là nội dung tiểu luận môn học Mạng cảm biến của riêng tôi và được thực hiện dưới sự hướng dẫn của giảng viên Hồ Nhựt Minh. Các phần trình bày trong báo cáo được xây dựng dựa trên quá trình tìm hiểu tài liệu, phân tích yêu cầu đề tài và thiết kế quy trình thực nghiệm trên nền tảng Edge Impulse.",
        "Trong phạm vi thực hiện, project Edge Impulse của đề tài đã upload dataset prototype 120 ảnh, generate features, huấn luyện FOMO, model testing, build deployment và chạy thử browser realtime trên Edge Impulse Studio. Các nội dung chưa có dữ liệu camera thật với chai nhựa ngoài hiện trường được đánh dấu rõ để bổ sung sau.",
        "Nếu phát hiện có bất kỳ sai lệch hoặc gian lận nào trong nội dung đã nộp, tôi xin hoàn toàn chịu trách nhiệm theo quy định của Học viện Công nghệ Bưu chính Viễn thông Cơ sở tại TP. Hồ Chí Minh.",
    ]
    for text in paragraphs:
        add_para(doc, text)
    add_para(doc, "TP. Hồ Chí Minh, ngày 05 tháng 06 năm 2026", align=WD_ALIGN_PARAGRAPH.RIGHT, first_line=False)
    add_para(doc, "Sinh viên", align=WD_ALIGN_PARAGRAPH.RIGHT, first_line=False, bold=True)
    add_para(doc, "\n\nPhạm Ngọc Quang Huy", align=WD_ALIGN_PARAGRAPH.RIGHT, first_line=False, bold=True)
    add_page_break(doc)

    add_heading(doc, "LỜI CẢM ƠN", 1)
    thanks = [
        "Em xin chân thành cảm ơn thầy Hồ Nhựt Minh đã giảng dạy học phần Mạng cảm biến và định hướng cách tiếp cận đề tài cuối kỳ. Thông qua học phần, em có cơ hội kết nối kiến thức cảm biến, truyền thông dữ liệu, xử lý tại biên và mô hình học máy nhúng thành một bài toán ứng dụng gần với thực tế.",
        "Em cũng xin cảm ơn Học viện Công nghệ Bưu chính Viễn thông Cơ sở tại TP. Hồ Chí Minh đã tạo điều kiện học tập, cung cấp môi trường thực hành và tài liệu tham khảo để sinh viên có thể triển khai các bài toán IoT/TinyML. Trong quá trình thực hiện, do thời gian và điều kiện thiết bị còn hạn chế, báo cáo không tránh khỏi thiếu sót; em mong nhận được góp ý của thầy để hoàn thiện hơn.",
    ]
    for text in thanks:
        add_para(doc, text)
    add_para(doc, "TP. Hồ Chí Minh, ngày 05 tháng 06 năm 2026", align=WD_ALIGN_PARAGRAPH.RIGHT, first_line=False)
    add_para(doc, "Sinh viên", align=WD_ALIGN_PARAGRAPH.RIGHT, first_line=False, bold=True)
    add_para(doc, "\n\nPhạm Ngọc Quang Huy", align=WD_ALIGN_PARAGRAPH.RIGHT, first_line=False, bold=True)
    add_page_break(doc)

    add_heading(doc, "TÓM TẮT", 1)
    summary = [
        "Bài tiểu luận trình bày đề tài phát hiện chai nhựa tái chế theo kích thước 500ml, 1L và 1.5L dựa trên nhãn dán trực quan và mô hình object detection FOMO của Edge Impulse. Bài toán hướng tới một nút cảm biến có camera đặt tại điểm thu gom hoặc băng chuyền phân loại rác, nơi hệ thống cần nhận diện nhanh loại chai để hỗ trợ thống kê và phân luồng tái chế.",
        "Phương pháp thực hiện gồm: khảo sát yêu cầu, thiết kế pipeline thu thập ảnh, gán nhãn bounding box với ba class 500ml, 1L, 1.5L, phân tích metric kích thước bbox_area_ratio, huấn luyện FOMO MobileNetV2 trên Edge Impulse, đánh giá bằng confusion matrix/F1 score và triển khai thử nghiệm WebAssembly trên trình duyệt. Do FOMO xuất centroid thay vì bounding box đầy đủ ở giai đoạn suy luận, metric kích thước bounding box được sử dụng như công cụ phân tích dữ liệu huấn luyện và kiểm tra chất lượng nhãn, không phải phép đo kích thước trực tiếp khi runtime.",
        "Kết quả thực nghiệm trên dataset prototype đạt validation F1 score 86.5%, test accuracy 86.67% và test F1 non-background 0.92. Hướng phát triển tiếp theo là bổ sung bộ dữ liệu thực tế tại nhiều điều kiện ánh sáng, tích hợp camera edge node, kiểm thử realtime bằng WebAssembly/browser và mở rộng hệ thống sang nhiều loại bao bì tái chế khác.",
    ]
    for text in summary:
        add_para(doc, text)


def front_matter_lists(doc: Document) -> None:
    add_page_break(doc)
    add_heading(doc, "MỤC LỤC", 1)
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    add_field(p, r'TOC \o "1-3" \h \z \u')
    add_para(doc, "Lưu ý: nếu mục lục chưa hiện số trang trong Word, chọn References > Update Table > Update entire table trước khi nộp.", italic=True, first_line=False, size=11)
    add_page_break(doc)

    add_heading(doc, "DANH MỤC CÁC THUẬT NGỮ, CHỮ VIẾT TẮT", 1)
    rows = [
        ("AI", "Artificial Intelligence", "Trí tuệ nhân tạo"),
        ("BBox", "Bounding Box", "Khung bao quanh đối tượng trong ảnh"),
        ("CNN", "Convolutional Neural Network", "Mạng nơ-ron tích chập"),
        ("Edge AI", "Edge Artificial Intelligence", "Trí tuệ nhân tạo chạy tại thiết bị biên"),
        ("FOMO", "Faster Objects, More Objects", "Mô hình object detection nhẹ của Edge Impulse"),
        ("IoT", "Internet of Things", "Internet vạn vật"),
        ("MCU", "Microcontroller Unit", "Vi điều khiển"),
        ("TinyML", "Tiny Machine Learning", "Học máy tối ưu cho thiết bị tài nguyên hạn chế"),
        ("WASM", "WebAssembly", "Định dạng triển khai chạy trong trình duyệt/JavaScript runtime"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for i, h in enumerate(["TỪ VIẾT TẮT", "TỪ GỐC", "Ý NGHĨA"]):
        table.cell(0, i).text = h
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val
    format_table(table, [3.0, 5.0, 8.0])
    add_page_break(doc)

    add_heading(doc, "DANH SÁCH BẢNG", 1)
    for item in [
        "Bảng 1.1. Đặc tả nhãn của đề tài",
        "Bảng 2.1. So sánh FOMO với mô hình object detection bounding box",
        "Bảng 3.1. Kế hoạch thu thập dữ liệu cho ba lớp chai",
        "Bảng 3.2. Metric bbox_area_ratio theo từng nhãn",
        "Bảng 4.1. Tổng hợp metric thực nghiệm từ Edge Impulse",
        "Bảng 4.2. Checklist minh chứng thực nghiệm",
        "Bảng 5.1. Tổng hợp kết quả đạt được và hướng phát triển",
    ]:
        add_para(doc, item, first_line=False)
    add_page_break(doc)

    add_heading(doc, "DANH SÁCH HÌNH", 1)
    for item in [
        "Hình 1.1. Kiến trúc tổng quát của hệ thống phát hiện chai nhựa tái chế",
        "Hình 2.1. Luồng suy luận FOMO và ý nghĩa centroid",
        "Hình 3.1. Giao diện upload dữ liệu trong project Edge Impulse của đề tài",
        "Hình 3.2. Data acquisition của project sau khi upload 120 mẫu chai nhựa",
        "Hình 3.3. Bounding-box size metric cho dữ liệu huấn luyện",
        "Hình 3.4. Feature generation cho tập training 90 mẫu và 3 class",
        "Hình 3.5. Pipeline Edge Impulse cho đề tài",
        "Hình 3.6. Cấu hình huấn luyện FOMO trong project của đề tài",
        "Hình 3.7. Deployment build thành công trong project Edge Impulse",
        "Hình 3.8. Browser realtime client đang chạy inferencing",
        "Hình 4.1. Dataset thật trong project Edge Impulse sau khi upload 120 mẫu chai nhựa",
        "Hình 4.2. Feature generation cho tập training 90 mẫu và 3 class",
        "Hình 4.3. Kết quả huấn luyện FOMO trên validation set",
        "Hình 4.4. Kết quả Model testing trên 30 mẫu test",
    ]:
        add_para(doc, item, first_line=False)


def chapter_1(doc: Document) -> None:
    add_heading(doc, "CHƯƠNG 1. TỔNG QUAN VỀ ĐỀ TÀI TIỂU LUẬN", 1)
    add_heading(doc, "1.1. Lý do chọn đề tài", 2)
    for text in [
        "Rác thải nhựa là một trong những nhóm vật liệu xuất hiện nhiều tại hộ gia đình, trường học, quán nước và điểm thu gom công cộng. Trong đó, chai nhựa dung tích 500ml, 1L và 1.5L rất phổ biến vì được dùng cho nước uống, trà, nước ngọt và các sản phẩm đóng chai. Khi các chai này được thu gom lẫn lộn, quá trình phân loại thủ công thường tốn thời gian, phụ thuộc vào người vận hành và khó tạo dữ liệu thống kê chính xác theo kích thước.",
        "Mạng cảm biến hiện đại không chỉ đo nhiệt độ, độ ẩm hay ánh sáng, mà còn có thể tích hợp camera và xử lý hình ảnh ngay tại thiết bị biên. Một node camera nhỏ có thể quan sát khu vực băng chuyền hoặc thùng thu gom, phát hiện chai theo nhãn dung tích, gửi kết quả về dashboard và hỗ trợ quyết định phân loại. Cách tiếp cận này phù hợp với xu hướng Edge AI, nơi dữ liệu được xử lý gần nguồn phát sinh nhằm giảm độ trễ, giảm băng thông và tăng tính riêng tư.",
        "Đề tài lựa chọn Edge Impulse và FOMO vì đây là nền tảng phù hợp cho sinh viên triển khai nhanh một bài toán TinyML/object detection. FOMO tập trung vào phát hiện vị trí đối tượng bằng centroid, nhẹ hơn nhiều so với các mô hình bounding box đầy đủ, nên có tiềm năng chạy trên thiết bị tài nguyên hạn chế. Với bài toán chai nhựa tái chế, hệ thống không nhất thiết phải đo chính xác toàn bộ khung bao ở runtime; yêu cầu chính là nhận diện đúng class chai và vị trí xuất hiện trong khung hình.",
    ]:
        add_para(doc, text)

    add_heading(doc, "1.2. Mục đích nghiên cứu", 2)
    for text in [
        "Mục đích của đề tài là xây dựng một quy trình phát hiện chai nhựa tái chế theo ba kích thước 500ml, 1L và 1.5L bằng mô hình FOMO trên Edge Impulse. Quy trình này bao gồm thiết kế dữ liệu, nhãn, pipeline huấn luyện, tiêu chí đánh giá và phương án triển khai thử nghiệm trên trình duyệt bằng WebAssembly.",
        "Bên cạnh nhận diện nhãn, đề tài đưa thêm metric bbox_area_ratio để đánh giá tính hợp lý của bounding box trong dữ liệu huấn luyện. Metric này không thay thế kết quả suy luận FOMO, mà giúp kiểm tra xem dữ liệu gán nhãn của class 500ml, 1L và 1.5L có phản ánh tương đối khác biệt kích thước hay không. Nếu metric của các lớp bị chồng lấn quá mạnh, sinh viên cần cải thiện cách chụp ảnh, khoảng cách camera hoặc quy tắc gán nhãn.",
    ]:
        add_para(doc, text)

    add_heading(doc, "1.3. Đối tượng và phạm vi nghiên cứu", 2)
    add_para(doc, "Đối tượng nghiên cứu của đề tài là ảnh chai nhựa tái chế có nhãn dán dung tích 500ml, 1L và 1.5L. Trong phạm vi tiểu luận, hệ thống tập trung vào phát hiện nhãn và vị trí chai trong ảnh, không xử lý nhận dạng thương hiệu, chất liệu nhựa chi tiết, mức độ bẩn của chai hoặc phân loại màu sắc nắp chai.")
    add_table_caption(doc, "Bảng 1.1. Đặc tả nhãn của đề tài")
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for i, h in enumerate(["Nhãn", "Đối tượng", "Ý nghĩa", "Ghi chú dữ liệu"]):
        table.cell(0, i).text = h
    rows = [
        ("500ml", "Chai nhựa dung tích khoảng 500ml", "Loại chai nhỏ, thường dùng cho nước uống cá nhân", "Chụp đủ góc đứng, nghiêng, xa/gần"),
        ("1L", "Chai nhựa dung tích khoảng 1 lít", "Loại trung bình, kích thước lớn hơn 500ml", "Cần tránh nhầm với 1.5L khi chai bị xa camera"),
        ("1.5L", "Chai nhựa dung tích khoảng 1.5 lít", "Loại lớn, thường dùng cho nước đóng chai gia đình", "Nên chụp cả ảnh có nhiều chai trong một khung"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val
    format_table(table, [2.5, 4.5, 4.6, 4.4])

    add_heading(doc, "1.4. Phương pháp nghiên cứu", 2)
    for item in [
        "Nghiên cứu tài liệu về mạng cảm biến, Edge AI, object detection, FOMO và triển khai WebAssembly trên Edge Impulse.",
        "Phân tích yêu cầu dữ liệu cho ba class chai nhựa và đề xuất quy tắc chụp ảnh, gán nhãn bounding box.",
        "Thiết kế impulse gồm khối xử lý ảnh và khối học FOMO MobileNetV2; xác định các metric cần theo dõi khi huấn luyện.",
        "Xây dựng metric bbox_area_ratio để rà soát chất lượng nhãn và mức độ tách biệt kích thước giữa ba class.",
        "Mô tả quy trình triển khai browser/WebAssembly và hướng tích hợp edge node trong hệ thống mạng cảm biến.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "1.5. Kiến trúc tổng quan của hệ thống", 2)
    add_para(doc, "Hệ thống đề xuất gồm bốn thành phần chính: camera thu ảnh, edge node xử lý cục bộ, mô hình FOMO đã huấn luyện và lớp hiển thị/ghi log kết quả. Camera quan sát vùng đặt chai hoặc băng chuyền. Edge node lấy từng frame, resize theo cấu hình impulse, chạy inference và gửi kết quả gồm nhãn, centroid, độ tin cậy cùng thời điểm phát hiện. Dashboard hoặc file log tổng hợp số lượng chai theo từng lớp để phục vụ thống kê tái chế.")
    arch = make_diagram_architecture()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.add_run().add_picture(str(arch), width=Cm(15.5))
    add_caption(doc, "Hình 1.1. Kiến trúc tổng quát của hệ thống phát hiện chai nhựa tái chế")


def chapter_2(doc: Document) -> None:
    add_heading(doc, "CHƯƠNG 2. CƠ SỞ LÝ THUYẾT", 1)
    add_heading(doc, "2.1. Mạng cảm biến có camera và xử lý tại biên", 2)
    for text in [
        "Mạng cảm biến là tập hợp các node có khả năng thu thập dữ liệu từ môi trường, xử lý cục bộ ở mức nhất định và truyền thông tin về hệ thống trung tâm. Trong các bài toán truyền thống, dữ liệu cảm biến thường là đại lượng vô hướng như nhiệt độ, độ ẩm, ánh sáng, rung động hoặc khí gas. Khi camera được đưa vào node cảm biến, dữ liệu trở thành ảnh hoặc video, có dung lượng lớn hơn và đòi hỏi phương pháp xử lý phức tạp hơn.",
        "Nếu mọi frame ảnh đều được gửi về server, hệ thống sẽ tiêu tốn băng thông, dễ tăng độ trễ và khó hoạt động ổn định trong môi trường mạng yếu. Edge AI giải quyết vấn đề này bằng cách đưa mô hình học máy về gần nguồn dữ liệu. Node chỉ gửi kết quả đã rút gọn như nhãn đối tượng, tọa độ, số lượng hoặc cảnh báo. Với đề tài này, edge node chỉ cần gửi thông tin rằng frame có chai 500ml, 1L hoặc 1.5L tại vị trí nào, thay vì gửi toàn bộ ảnh liên tục.",
        "Một lợi ích khác của xử lý tại biên là khả năng phản hồi nhanh. Trong trạm phân loại rác, việc phát hiện sai hoặc chậm có thể khiến chai đi qua vị trí phân loại. Khi inference chạy cục bộ, độ trễ giảm và hệ thống có thể đưa ra tín hiệu điều khiển cơ cấu phân luồng hoặc cảnh báo gần thời gian thực. Đây là lý do đề tài ưu tiên mô hình nhẹ như FOMO thay vì mô hình lớn cần GPU.",
    ]:
        add_para(doc, text)

    add_heading(doc, "2.2. Object detection và bounding box", 2)
    for text in [
        "Object detection là bài toán nhận biết lớp đối tượng và vị trí của đối tượng trong ảnh. Với mô hình bounding box truyền thống, đầu ra thường gồm nhãn, độ tin cậy và khung bao có tọa độ x, y, width, height. Khung bao này phù hợp khi cần đo kích thước, cắt vùng ảnh hoặc theo dõi đối tượng bằng IoU giữa các frame.",
        "Trong quá trình xây dựng dataset, bounding box vẫn rất quan trọng dù mô hình runtime có xuất bbox hay không. Bounding box giúp người gán nhãn chỉ rõ vùng chứa chai, tránh để mô hình học nhiễu từ nền. Các thông tin width và height của bbox cũng có thể dùng để phân tích chất lượng dataset, ví dụ kiểm tra class 1.5L có xu hướng chiếm diện tích ảnh lớn hơn class 500ml khi camera đặt cùng khoảng cách hay không.",
    ]:
        add_para(doc, text)

    add_heading(doc, "2.3. FOMO trong Edge Impulse", 2)
    for text in [
        "FOMO, viết tắt của Faster Objects, More Objects, là learning block object detection nhẹ của Edge Impulse. Thay vì dự đoán bounding box đầy đủ như nhiều mô hình object detection cổ điển, FOMO tạo heatmap theo lưới và trả về vị trí tâm của đối tượng. Cách tiếp cận này làm giảm chi phí tính toán, phù hợp với các thiết bị biên có RAM và CPU hạn chế.",
        "Tài liệu Edge Impulse nhấn mạnh giới hạn quan trọng của FOMO: mô hình không xuất bounding box đầy đủ ở inference mà xuất centroid, vì vậy kích thước đối tượng không có sẵn trực tiếp. Đồng thời, FOMO hoạt động tốt hơn khi các đối tượng có kích thước tương đối giống nhau và không đứng quá sát nhau trong ảnh. Với đề tài chai nhựa, điều này cần được xử lý bằng cách chuẩn hóa góc đặt camera, khoảng cách chụp và quy tắc bố trí chai trong vùng quan sát.",
        "FOMO sử dụng MobileNetV2 làm backbone và áp dụng classifier theo dạng fully convolutional trên bản đồ đặc trưng. Với cấu hình thường dùng, ảnh đầu vào vuông như 96x96 có thể tạo output grid 12x12, mỗi ô tương ứng một vùng trong ảnh. Khi heatmap có activation đủ mạnh cho class 500ml, 1L hoặc 1.5L, hệ thống quy đổi ô đó thành vị trí centroid của đối tượng.",
    ]:
        add_para(doc, text)

    inf = make_diagram_inference()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.add_run().add_picture(str(inf), width=Cm(15.5))
    add_caption(doc, "Hình 2.1. Luồng suy luận FOMO và ý nghĩa centroid")

    add_heading(doc, "2.4. So sánh FOMO với mô hình bounding box", 2)
    add_table_caption(doc, "Bảng 2.1. So sánh FOMO với mô hình object detection bounding box")
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for i, h in enumerate(["Tiêu chí", "FOMO", "Mô hình bbox đầy đủ", "Ý nghĩa với đề tài"]):
        table.cell(0, i).text = h
    rows = [
        ("Đầu ra", "Nhãn, centroid, score", "Nhãn, bbox, score", "FOMO phù hợp nếu cần vị trí và class hơn là kích thước runtime"),
        ("Tài nguyên", "Nhẹ, hướng tới MCU/edge", "Thường nặng hơn", "Dễ triển khai trên node cảm biến giá rẻ"),
        ("Dữ liệu gán nhãn", "Dùng bounding box khi training", "Dùng bounding box khi training", "Vẫn cần gán nhãn cẩn thận cho 500ml/1L/1.5L"),
        ("Kích thước đối tượng", "Không có sẵn trực tiếp ở inference", "Có width/height của bbox", "Metric bbox_area_ratio chỉ dùng để phân tích dữ liệu"),
        ("Khoảng cách đối tượng", "Nhạy khi vật quá sát nhau", "Thường xử lý tốt hơn", "Cần bố trí camera/băng chuyền tránh chồng lấn"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val
    format_table(table, [3.0, 3.7, 3.9, 5.4])

    add_heading(doc, "2.5. WebAssembly trong kiểm thử trình duyệt", 2)
    add_para(doc, "Edge Impulse cho phép đóng gói impulse thành WebAssembly để chạy trực tiếp trong trình duyệt hoặc môi trường JavaScript. Với bài tiểu luận, hướng triển khai browser có hai lợi ích: sinh viên dễ quay/chụp minh chứng realtime bằng webcam, và không cần chuẩn bị ngay phần cứng nhúng. Sau khi mô hình đạt kết quả chấp nhận được, cùng impulse có thể tiếp tục được xuất sang C++ library, Arduino library hoặc Linux .eim để tích hợp vào thiết bị thật.")

    add_heading(doc, "2.6. Ràng buộc TinyML trong node cảm biến", 2)
    for text in [
        "Khi đưa mô hình học máy vào node cảm biến, tài nguyên tính toán là ràng buộc đầu tiên cần xem xét. Một camera node giá rẻ thường có CPU yếu, RAM hạn chế, bộ nhớ flash nhỏ và nguồn cấp không ổn định như máy tính. Nếu mô hình quá lớn, thiết bị có thể không nạp được model, inference chậm hoặc treo khi xử lý liên tục. Vì vậy, trong học phần Mạng cảm biến, lựa chọn mô hình không chỉ dựa trên độ chính xác mà còn dựa trên khả năng triển khai thực tế.",
        "TinyML là hướng tối ưu mô hình để chạy trên vi điều khiển hoặc thiết bị biên nhỏ. Các kỹ thuật thường gặp gồm giảm kích thước ảnh đầu vào, lượng tử hóa int8, dùng backbone nhẹ, giảm số lớp, giảm số tham số và chỉ gửi kết quả sau inference thay vì gửi raw image. FOMO phù hợp với triết lý này vì mô hình ưu tiên vị trí centroid và class, giảm gánh nặng tính toán so với mô hình bounding box đầy đủ.",
        "Với đề tài chai nhựa, việc giảm input về ảnh vuông như 96x96 hoặc 160x160 có thể làm mất một phần chi tiết chữ trên nhãn dán. Do đó, cần cân bằng giữa tốc độ và khả năng nhận diện nhãn. Nếu ảnh 96x96 làm mô hình nhầm 1L với 1.5L, có thể tăng kích thước ảnh hoặc chụp nhãn rõ hơn thay vì chỉ tăng epoch huấn luyện. Một mô hình nhanh nhưng học từ dữ liệu kém vẫn cho kết quả không ổn định.",
        "Ngoài mô hình, node cảm biến còn phải xử lý pipeline đầy đủ: lấy frame, resize, chuẩn hóa màu, chạy inference, lọc threshold, gửi dữ liệu và đôi khi hiển thị kết quả. Các bước này đều tiêu tốn thời gian. Khi thiết kế thực tế, sinh viên nên đo latency tổng thể theo frame, không chỉ xem thời gian inference do Edge Impulse báo. Latency tổng thể mới phản ánh khả năng dùng mô hình trên băng chuyền hoặc điểm thu gom rác.",
    ]:
        add_para(doc, text)

    add_heading(doc, "2.7. Các metric đánh giá mô hình", 2)
    for text in [
        "Đối với object detection, đánh giá mô hình cần quan tâm cả nhận diện class và định vị đối tượng. Với FOMO, định vị được biểu diễn bằng centroid nên việc đánh giá vị trí có thể dựa trên khoảng cách giữa centroid dự đoán và tâm của bounding box gán nhãn. Nếu centroid nằm gần tâm chai, hệ thống có thể dùng vị trí này để đếm hoặc điều khiển cơ cấu phân loại. Nếu centroid lệch ra ngoài thân chai, phát hiện đó vẫn cần xem là chưa tốt dù nhãn class đúng.",
        "Precision cho biết trong các lần mô hình báo có chai, bao nhiêu lần là đúng. Recall cho biết trong các chai thật có trong ảnh, mô hình phát hiện được bao nhiêu. F1 score là trung bình điều hòa giữa precision và recall, thường được dùng khi cần cân bằng false positive và false negative. Trong bài toán phân loại rác, false positive có thể khiến hệ thống đếm dư hoặc kích hoạt sai, còn false negative làm bỏ sót chai cần thu gom.",
        "Confusion matrix giúp thấy cặp nhãn nào hay bị nhầm. Nếu 500ml hiếm khi nhầm với 1.5L nhưng 1L và 1.5L nhầm nhiều, vấn đề có thể nằm ở góc chụp, khoảng cách camera hoặc kiểu dáng chai giống nhau. Khi có confusion matrix thật, sinh viên cần thảo luận từng lỗi nhầm chính thay vì chỉ ghi một con số accuracy tổng. Điều này làm báo cáo có giá trị kỹ thuật hơn và giúp đề xuất cải thiện dữ liệu cụ thể.",
        "Bên cạnh metric của Edge Impulse, bbox_area_ratio đóng vai trò kiểm tra dữ liệu. Metric này không phải metric accuracy của FOMO, nhưng giúp phát hiện ảnh chụp không đồng nhất. Ví dụ nếu class 500ml có nhiều ảnh được chụp sát camera, phân bố diện tích bbox có thể lớn hơn class 1.5L. Khi đó, mô hình có thể học sai tương quan giữa kích thước trên ảnh và nhãn dung tích. Phát hiện sớm lỗi này giúp giảm thời gian huấn luyện lặp lại.",
    ]:
        add_para(doc, text)


def chapter_3(doc: Document) -> None:
    add_heading(doc, "CHƯƠNG 3. THIẾT KẾ VÀ PHƯƠNG PHÁP THỰC HIỆN", 1)
    add_heading(doc, "3.1. Mô tả bài toán", 2)
    for text in [
        "Bài toán đầu vào là ảnh RGB hoặc grayscale chứa một hoặc nhiều chai nhựa tái chế. Mỗi chai cần được gán một trong ba nhãn: 500ml, 1L hoặc 1.5L. Đầu ra mong muốn của hệ thống khi chạy FOMO là danh sách phát hiện gồm class, centroid và độ tin cậy. Nếu một frame có nhiều chai, hệ thống có thể trả về nhiều centroid tương ứng.",
        "Trong bối cảnh phân loại rác, nhãn dung tích được xem như dấu hiệu phân loại chính. Ở giai đoạn dữ liệu, người thực hiện cần chụp rõ hình dáng chai và nhãn dán dung tích. Các chai bị che khuất nhiều, mất nhãn, biến dạng quá mạnh hoặc nằm ngoài vùng quan sát nên được loại khỏi tập huấn luyện ban đầu để tránh làm nhiễu mô hình.",
    ]:
        add_para(doc, text)

    add_heading(doc, "3.2. Project Edge Impulse của đề tài", 2)
    add_para(doc, "Đề tài được triển khai trong project Edge Impulse tên Recycling bottle FOMO 500ml 1L 1.5L. Project sử dụng ba class đầu ra đúng theo yêu cầu: 500ml, 1L và 1.5L. Dữ liệu trong project là dataset prototype 120 mẫu chai nhựa có bounding box nhất quán, được dùng để kiểm chứng đầy đủ pipeline từ upload dữ liệu, generate features, huấn luyện FOMO, model testing đến deployment/browser realtime.")
    if add_screenshot(doc, "07_cloned_project_upload_modal.png", "Hình 3.1. Giao diện upload dữ liệu trong project Edge Impulse của đề tài"):
        add_para(doc, "Screenshot thể hiện project của đề tài với tên Recycling bottle FOMO 500ml 1L 1.5L và màn hình upload dữ liệu object detection có hỗ trợ annotation bounding box.", italic=True, first_line=False, size=11)
    else:
        add_placeholder(doc, "PLACEHOLDER HÌNH 3.1 - PROJECT EDGE IMPULSE", "Chèn screenshot project Edge Impulse của đề tài 500ml/1L/1.5L")
        add_caption(doc, "Hình 3.1. Placeholder minh chứng project Edge Impulse của đề tài")

    add_heading(doc, "3.3. Kế hoạch thu thập dữ liệu", 2)
    add_para(doc, "Dataset cần phản ánh điều kiện sử dụng thực tế. Mỗi class nên có ảnh ở nhiều khoảng cách, nhiều góc xoay, ánh sáng khác nhau và cả tình huống có nhiều chai trong một frame. Để tránh mô hình chỉ học nền, cần thay đổi mặt bàn, thùng rác, băng chuyền giả lập hoặc phông nền khi chụp.")
    if add_screenshot(doc, "08_real_uploaded_bottle_dataset.png", "Hình 3.2. Data acquisition của project sau khi upload 120 mẫu chai nhựa"):
        add_para(doc, "Screenshot là minh chứng thật từ project của đề tài: dataset có 120 items, train/test split 75%/25%, gồm 90 mẫu training và 30 mẫu testing với các nhãn 500ml, 1L và 1.5L.", italic=True, first_line=False, size=11)
    add_table_caption(doc, "Bảng 3.1. Kế hoạch thu thập dữ liệu cho ba lớp chai")
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    for i, h in enumerate(["Class", "Số ảnh khuyến nghị", "Biến thiên cần có", "Tỷ lệ train/test", "Ghi chú"]):
        table.cell(0, i).text = h
    rows = [
        ("500ml", ">= 150 ảnh", "Đứng, nghiêng, xa/gần, nền sáng/tối", "80/20", "Tránh chỉ chụp một thương hiệu"),
        ("1L", ">= 150 ảnh", "Đứng, nghiêng, có/không móp nhẹ", "80/20", "Giữ cùng khoảng cách camera với 500ml"),
        ("1.5L", ">= 150 ảnh", "Chai đơn và nhiều chai trong frame", "80/20", "Bổ sung ảnh xa để giảm lệch kích thước tuyệt đối"),
        ("Background", ">= 50 ảnh", "Không có chai hoặc vật gây nhiễu", "80/20", "Giúp giảm false positive"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val
    format_table(table, [2.2, 3.0, 4.4, 2.2, 4.2])

    add_heading(doc, "3.4. Quy tắc gán nhãn bounding box", 2)
    for item in [
        "Bounding box bao sát thân chai và nhãn dung tích, không lấy quá nhiều nền xung quanh.",
        "Nếu chai bị nghiêng, vẫn vẽ bbox trục ngang/dọc chứa toàn bộ phần chai nhìn thấy.",
        "Không gán nhãn cho vật thể bị che khuất quá nhiều hoặc không xác định được dung tích.",
        "Nếu ảnh có nhiều chai, mỗi chai là một bbox riêng với nhãn tương ứng.",
        "Giữ tên class thống nhất tuyệt đối: 500ml, 1L, 1.5L; không dùng biến thể như 500 ml hoặc 1 lít.",
    ]:
        add_bullet(doc, item)
    if add_screenshot(doc, "08_real_uploaded_bottle_dataset.png", "Hình 3.2a. Dataset prototype đã upload với 120 mẫu và ba nhãn 500ml, 1L, 1.5L"):
        add_para(doc, "Screenshot này là minh chứng thật từ project Edge Impulse của đề tài: 120 mẫu chai nhựa prototype, gồm 90 mẫu training và 30 mẫu testing, mỗi mẫu có bounding box và nhãn đúng theo ba class 500ml, 1L và 1.5L.", italic=True, first_line=False, size=11)
    else:
        add_placeholder(doc, "PLACEHOLDER - LABELING QUEUE 500ml/1L/1.5L", "Cần chèn screenshot thật của project riêng khi vẽ bounding box và chọn nhãn 500ml, 1L, 1.5L.")

    add_heading(doc, "3.5. Bounding-box size metric", 2)
    add_para(doc, "Metric bbox_area_ratio được đề xuất để kiểm tra chất lượng dữ liệu và sự khác biệt kích thước tương đối giữa các class. Với mỗi bounding box, diện tích bbox được chia cho diện tích ảnh. Công thức sử dụng:")
    add_para(doc, "bbox_area_ratio = (bbox_width x bbox_height) / (image_width x image_height)", align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False, bold=True)
    add_para(doc, "Nếu camera đặt ổn định, class 1.5L thường có bbox_area_ratio lớn hơn class 1L và 500ml. Tuy nhiên, metric này phụ thuộc mạnh vào khoảng cách camera. Vì vậy, metric không được dùng một mình để kết luận dung tích chai, mà dùng để phát hiện dữ liệu bất thường, ví dụ ảnh 500ml chụp quá gần khiến bbox lớn hơn ảnh 1.5L chụp quá xa.")
    metric = make_diagram_metric()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.add_run().add_picture(str(metric), width=Cm(15.0))
    add_caption(doc, "Hình 3.3. Bounding-box size metric cho dữ liệu huấn luyện")
    add_table_caption(doc, "Bảng 3.2. Metric bbox_area_ratio theo từng nhãn")
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    for i, h in enumerate(["Class", "Mean", "Median", "IQR", "Diễn giải"]):
        table.cell(0, i).text = h
    for row in [
        ("500ml", "0.06330", "0.06349", "0.00809", "Thấp nhất trong dataset prototype đã upload"),
        ("1L", "0.09584", "0.09449", "0.00736", "Nằm giữa 500ml và 1.5L"),
        ("1.5L", "0.13828", "0.13838", "0.01119", "Cao nhất trong dataset prototype đã upload"),
    ]:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val
    format_table(table, [2.2, 2.2, 2.2, 2.2, 7.2])

    add_heading(doc, "3.6. Thiết kế impulse", 2)
    for text in [
        "Impulse gồm khối đầu vào ảnh, khối xử lý ảnh và learning block Object Detection (Images) với mô hình FOMO. Kích thước ảnh có thể bắt đầu từ 96x96 để ưu tiên tốc độ, sau đó tăng lên 160x160 nếu centroid quá lệch hoặc chai nhỏ khó phát hiện. Resize mode nên dùng Fit shortest axis để giữ tỷ lệ khung hình trước khi crop về ảnh vuông.",
        "Ở khối Image, sinh viên có thể thử RGB trước vì nhãn dán và màu chai có thể hỗ trợ nhận diện. Nếu thiết bị biên quá hạn chế RAM, grayscale là phương án giảm tài nguyên. Với bài báo cáo, lựa chọn mặc định là RGB để giữ nhiều thông tin hình ảnh trong giai đoạn thử nghiệm trên trình duyệt.",
    ]:
        add_para(doc, text)
    if add_screenshot(doc, "09_real_feature_generation.png", "Hình 3.4. Feature generation cho tập training 90 mẫu và 3 class"):
        add_para(doc, "Feature explorer trong project hiển thị đủ ba class 1.5L, 1L và 500ml. Khối xử lý ảnh báo processing time 7 ms và peak RAM 4 KB trước learning block.", italic=True, first_line=False, size=11)
    pipeline = make_diagram_pipeline()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.add_run().add_picture(str(pipeline), width=Cm(15.5))
    add_caption(doc, "Hình 3.5. Pipeline Edge Impulse cho đề tài")

    add_heading(doc, "3.7. Huấn luyện và đánh giá", 2)
    for text in [
        "Quá trình huấn luyện được thực hiện trong Edge Impulse Studio. Các tham số như số epoch, learning rate và augmentation nên bắt đầu theo template rồi điều chỉnh dần. Augmentation nhẹ như thay đổi brightness, crop nhỏ hoặc flip phù hợp có thể giúp mô hình bền hơn với điều kiện thực tế. Tuy nhiên, augmentation quá mạnh làm biến dạng nhãn dung tích hoặc hình dáng chai cần tránh.",
        "Các metric đánh giá chính gồm confusion matrix, F1 score, precision/recall theo từng class và kết quả kiểm thử bằng ảnh chưa dùng trong training. Với FOMO, cần kiểm tra thêm vị trí centroid có nằm gần tâm chai hay không. Nếu mô hình nhận đúng class nhưng centroid lệch nhiều, khi đưa vào băng chuyền hoặc cơ cấu phân loại vẫn có thể gây lỗi định vị.",
    ]:
        add_para(doc, text)
    if add_screenshot(doc, "10_before_fomo_training_real.png", "Hình 3.6. Cấu hình huấn luyện FOMO trong project của đề tài"):
        add_para(doc, "Screenshot thể hiện cấu hình huấn luyện trong project của đề tài: 40 training cycles, learning rate 0.001, training processor CPU và bật data augmentation.", italic=True, first_line=False, size=11)
    if add_screenshot(doc, "11_real_fomo_training_results.png", "Hình 3.6b. Kết quả huấn luyện FOMO thật cho dataset 500ml, 1L, 1.5L"):
        add_para(doc, "Kết quả thật trên project của đề tài: FOMO MobileNetV2 0.35, output 3 classes, validation F1 score 86.5%, precision 0.84, recall 0.89, F1 non-background 0.86; inferencing time 630 ms, peak RAM 114.7K và flash 81.0K trên target Cortex-M4F 80MHz.", italic=True, first_line=False, size=11)
    else:
        add_placeholder(doc, "PLACEHOLDER - TRAINING RESULT 500ml/1L/1.5L", "Chèn screenshot kết quả training, confusion matrix và F1 score thật từ Edge Impulse sau khi huấn luyện dataset chai nhựa.")

    add_heading(doc, "3.8. Triển khai WebAssembly/browser realtime", 2)
    add_para(doc, "Sau khi mô hình đạt kết quả chấp nhận được, project được build trong mục Deployment của Edge Impulse và chạy thử bằng browser realtime client. Khi test realtime, webcam gửi từng frame cho runtime trình duyệt, mô hình chạy cục bộ và trả về trạng thái inferencing theo thời gian thực. Do lượt demo này không đặt chai nhựa thật trước webcam, screenshot realtime chỉ chứng minh runtime browser đã chạy, chưa phải minh chứng nhận diện chai ngoài hiện trường.")
    if add_screenshot(doc, "14_real_deployment_build_success.png", "Hình 3.7. Deployment build thành công trong project Edge Impulse"):
        add_para(doc, "Build deployment thật đã hoàn tất với phiên bản v12 (C++ library), thời điểm 21:40:07 ngày 05/06/2026. Trang deployment cũng hiển thị Launch in browser để chạy prototype realtime.", italic=True, first_line=False, size=11)
    if add_screenshot(doc, "15_real_browser_realtime_inferencing_full_redacted.png", "Hình 3.8. Browser realtime client đang chạy inferencing"):
        add_para(doc, "Browser realtime client đã tải model và chuyển sang trạng thái Inferencing với webcam. Vùng camera trong ảnh được làm mờ để bảo vệ riêng tư; do không có chai nhựa trong khung hình, báo cáo không ghi nhận nhãn phát hiện giả.", italic=True, first_line=False, size=11)

    add_heading(doc, "3.9. Thiết kế tích hợp phần cứng đề xuất", 2)
    for text in [
        "Trong phiên bản triển khai thật, edge node có thể là Raspberry Pi có camera USB/CSI hoặc một board MCU có camera tùy điều kiện thiết bị. Raspberry Pi thuận lợi cho thử nghiệm vì có hệ điều hành Linux, dễ chạy Edge Impulse Linux SDK hoặc WebAssembly thông qua môi trường browser/Node.js. ESP32-CAM có chi phí thấp hơn nhưng tài nguyên hạn chế hơn, phù hợp khi mô hình đã được tối ưu kỹ và chỉ cần gửi kết quả đơn giản.",
        "Camera nên được cố định ở một vị trí có khoảng cách tương đối ổn định tới vùng đặt chai. Nếu camera đặt quá thấp, chai cao 1.5L có thể bị cắt đầu; nếu đặt quá xa, nhãn 500ml dễ nhỏ và khó nhận diện. Góc chụp nên hạn chế phản quang trên thân chai trong suốt. Trong môi trường có băng chuyền, nên bố trí ánh sáng tán xạ để giảm bóng đổ và giữ tốc độ băng chuyền phù hợp với tốc độ inference.",
        "Kết quả inference có thể gửi qua MQTT, HTTP hoặc WebSocket tùy kiến trúc hệ thống. Gói tin tối thiểu gồm timestamp, label, x, y, score và device_id. Nếu cần ghi ảnh để audit, chỉ lưu ảnh khi score thấp hoặc khi người vận hành xác nhận mô hình sai, tránh lưu mọi frame gây đầy bộ nhớ. Cách này giữ đúng tinh thần mạng cảm biến: node xử lý cục bộ và chỉ truyền thông tin có giá trị.",
        "Nguồn cấp và độ ổn định mạng cũng cần được tính đến. Nếu node mất mạng, hệ thống nên lưu log tạm cục bộ rồi đồng bộ lại khi kết nối trở lại. Nếu camera hoạt động ngoài trời hoặc gần khu rác, hộp bảo vệ cần chống bụi và giữ ống kính sạch. Những yếu tố này không thuộc trực tiếp mô hình FOMO nhưng quyết định khả năng vận hành của hệ thống trong thực tế.",
    ]:
        add_para(doc, text)

    add_heading(doc, "3.10. Kiểm soát chất lượng dữ liệu", 2)
    for text in [
        "Chất lượng dữ liệu quyết định phần lớn kết quả của mô hình. Với ba nhãn 500ml, 1L và 1.5L, sinh viên cần tránh dataset bị lệch, ví dụ 500ml chỉ chụp trên nền trắng còn 1.5L chỉ chụp trên nền tối. Khi nền gắn chặt với class, mô hình có thể học nền thay vì học chai. Cách kiểm tra đơn giản là trộn ngẫu nhiên ảnh theo class và quan sát bằng mắt xem có pattern nền quá rõ hay không.",
        "Mỗi class nên có sự đa dạng về thương hiệu, độ trong của nhựa, mức nước còn lại, độ móp và hướng nhãn. Tuy nhiên, trong giai đoạn đầu không nên đưa quá nhiều biến thể khó cùng lúc. Nên bắt đầu bằng ảnh rõ, chai nằm trong vùng quan sát và nhãn dung tích đọc được, sau đó tăng dần độ khó bằng ảnh nghiêng, chai bị móp nhẹ hoặc ánh sáng phức tạp. Chiến lược này giúp dễ xác định nguyên nhân khi mô hình lỗi.",
        "Sau khi gán nhãn, cần kiểm tra các bbox có quá rộng hoặc quá hẹp. Bbox quá rộng làm mô hình học nhiều nền, còn bbox quá hẹp có thể bỏ mất phần đặc trưng của thân chai. Với FOMO, tâm của bbox cũng quan trọng vì mô hình học centroid. Nếu bbox lệch khỏi thân chai, centroid mục tiêu bị lệch và mô hình runtime dễ trả vị trí không đúng.",
        "Một bước nên làm là xuất danh sách nhãn ra file JSON hoặc CSV rồi tính thống kê số lượng mẫu, bbox_area_ratio, width/height ratio theo class. Các ảnh ngoại lai có bbox_area_ratio quá nhỏ/quá lớn nên được mở lại để kiểm tra. Nếu ngoại lai là ảnh hợp lệ do khoảng cách chụp khác, nên cân nhắc tách thành nhóm test hoặc chụp bổ sung để cân bằng phân bố.",
    ]:
        add_para(doc, text)


def chapter_4(doc: Document) -> None:
    add_heading(doc, "CHƯƠNG 4. KẾT QUẢ VÀ THẢO LUẬN", 1)
    add_heading(doc, "4.1. Kết quả thiết kế đạt được", 2)
    for text in [
        "Trong phạm vi báo cáo hiện tại, đề tài đã hoàn thành thiết kế quy trình triển khai FOMO cho bài toán phát hiện chai nhựa theo ba nhãn dung tích. Báo cáo đã xác định rõ class, quy tắc gán nhãn, metric kiểm tra kích thước bbox, kiến trúc hệ thống mạng cảm biến và phương án triển khai WebAssembly.",
        "Trong project Edge Impulse của đề tài, dataset prototype 120 ảnh chai nhựa có nhãn 500ml, 1L và 1.5L đã được upload và chia thành 90 mẫu training, 30 mẫu testing. Mô hình FOMO đã được huấn luyện, kiểm thử, build deployment và chạy thử browser realtime trên Edge Impulse.",
    ]:
        add_para(doc, text)

    add_heading(doc, "4.2. Kết quả thực nghiệm trên Edge Impulse", 2)
    add_para(doc, "Dataset thực nghiệm trong lượt triển khai này là dataset prototype được sinh tự động để có nhãn và bounding box nhất quán, không phải ảnh thu thập ngoài hiện trường. Mục tiêu của bước này là chạy thật pipeline Edge Impulse từ upload dữ liệu, generate features, train FOMO đến Model testing, đồng thời không bịa số liệu khi chưa có dữ liệu camera thực tế.")
    if add_screenshot(doc, "08_real_uploaded_bottle_dataset.png", "Hình 4.1. Dataset thật trong project Edge Impulse sau khi upload 120 mẫu chai nhựa"):
        add_para(doc, "Project hiển thị 120 items, train/test split 75%/25%, gồm 90 mẫu training và 30 mẫu testing. Các sample trong danh sách có nhãn 1.5L, 1L và 500ml đúng theo ba class của đề tài.", italic=True, first_line=False, size=11)
    if add_screenshot(doc, "09_real_feature_generation.png", "Hình 4.2. Feature generation cho tập training 90 mẫu và 3 class"):
        add_para(doc, "Feature explorer hiển thị đủ ba class 1.5L, 1L và 500ml. Khối xử lý ảnh báo processing time 7 ms và peak RAM 4 KB trước learning block.", italic=True, first_line=False, size=11)
    if add_screenshot(doc, "11_real_fomo_training_results.png", "Hình 4.3. Kết quả huấn luyện FOMO trên validation set"):
        add_para(doc, "Kết quả validation: F1 score 86.5%, precision non-background 0.84, recall non-background 0.89, F1 non-background 0.86. Confusion matrix cho thấy class 1L đạt 100%, class 1.5L đạt 83.3% và class 500ml đạt 87.5% theo hàng class thật.", italic=True, first_line=False, size=11)
    if add_screenshot(doc, "12_real_model_testing_results.png", "Hình 4.4. Kết quả Model testing trên 30 mẫu test"):
        add_para(doc, "Kết quả test set: accuracy 86.67%, precision non-background 0.90, recall non-background 0.93 và F1 non-background 0.92. Một số mẫu test 1L có F1 thấp, cho thấy cần bổ sung ảnh thật đa dạng hơn nếu muốn nghiệm thu ngoài môi trường thực tế.", italic=True, first_line=False, size=11)

    add_table_caption(doc, "Bảng 4.1. Tổng hợp metric thực nghiệm từ Edge Impulse")
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for i, h in enumerate(["Hạng mục", "Tập dữ liệu", "Giá trị", "Ghi chú"]):
        table.cell(0, i).text = h
    rows = [
        ("Số mẫu", "Training/Test", "90 / 30", "Tổng 120 ảnh prototype có bbox"),
        ("Validation F1", "Validation set", "86.5%", "FOMO MobileNetV2 0.35, 40 cycles"),
        ("Validation precision/recall", "Validation set", "0.84 / 0.89", "Non-background metrics"),
        ("Test accuracy", "Test set", "86.67%", "Chạy bằng Classify all"),
        ("Test precision/recall/F1", "Test set", "0.90 / 0.93 / 0.92", "Non-background metrics"),
        ("On-device estimate", "Cortex-M4F 80MHz", "630 ms, 114.7K RAM, 81.0K flash", "Theo trang Object detection"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val
    format_table(table, [3.6, 3.0, 4.1, 5.3])

    add_page_break(doc)
    add_heading(doc, "4.3. Checklist minh chứng thực nghiệm", 2)
    add_table_caption(doc, "Bảng 4.2. Checklist minh chứng thực nghiệm")
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for i, h in enumerate(["STT", "Minh chứng", "Trạng thái trong báo cáo", "Ghi chú thay thế trước khi nộp"]):
        table.cell(0, i).text = h
    rows = [
        ("1", "Project Edge Impulse", "Đã có minh chứng", "Project Recycling bottle FOMO 500ml 1L 1.5L"),
        ("2", "Data acquisition", "Đã có minh chứng thật", "120 ảnh prototype, 90 training, 30 testing"),
        ("3", "Labeling/bounding boxes", "Đã có minh chứng thật", "Mỗi sample có bbox và nhãn 500ml/1L/1.5L"),
        ("4", "Training FOMO", "Đã có minh chứng thật", "Validation F1 86.5%, confusion matrix 3 class"),
        ("5", "Model testing", "Đã có minh chứng thật", "Test accuracy 86.67%, F1 non-background 0.92"),
        ("6", "Deployment build", "Đã có minh chứng thật", "Build v12 C++ library lúc 21:40:07 ngày 05/06/2026"),
        ("7", "Realtime browser test", "Đã có minh chứng thật", "Browser client đã chạy Inferencing; camera frame được làm mờ để bảo vệ riêng tư"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val
    format_table(table, [1.2, 4.5, 3.0, 7.3])

    add_heading(doc, "4.4. Thảo luận về FOMO và kích thước chai", 2)
    for text in [
        "Điểm mạnh của FOMO là nhẹ, dễ triển khai tại biên và phù hợp với bài toán cần biết đối tượng xuất hiện ở đâu trong ảnh. Với chai nhựa trên băng chuyền, centroid đủ để xác định vùng tác động của cơ cấu phân loại hoặc vị trí cần ghi log. Mô hình cũng có thể phát hiện nhiều đối tượng nếu chúng cách nhau đủ rõ trong ảnh.",
        "Điểm cần lưu ý là FOMO không cung cấp kích thước bounding box khi suy luận. Vì vậy, nếu hệ thống tương lai cần đo kích thước chai trực tiếp trong runtime, có hai hướng mở rộng: chuyển sang mô hình bounding box như MobileNetV2 SSD FPN/YOLO, hoặc thêm bước hậu xử lý thị giác máy tính dựa trên segmentation/contour quanh centroid. Trong phiên bản đề tài này, nhãn dung tích được xem là class phân loại, còn bbox size metric chỉ là công cụ phân tích dataset.",
        "Metric bbox_area_ratio vẫn có giá trị trong quá trình chuẩn bị dữ liệu. Khi dữ liệu được chụp ở khoảng cách tương đối ổn định, phân bố bbox_area_ratio của 500ml, 1L và 1.5L sẽ cho thấy việc gán nhãn có hợp lý hay không. Nếu class 1.5L có nhiều bbox nhỏ bất thường, có thể ảnh đó chụp quá xa hoặc bbox vẽ thiếu thân chai. Nếu class 500ml có bbox quá lớn, có thể ảnh chụp quá gần và gây lệch phân bố.",
    ]:
        add_para(doc, text)

    add_heading(doc, "4.5. Các lỗi có thể gặp", 2)
    for item in [
        "Nhầm lẫn giữa 1L và 1.5L khi chai bị chụp xa, mất nhãn hoặc có kiểu dáng tương tự.",
        "False positive với chai thủy tinh, lon nước hoặc vật có hình trụ nếu dataset background quá ít.",
        "Centroid lệch khi chai nằm sát mép ảnh hoặc bị che khuất bởi chai khác.",
        "Mô hình học theo nền nếu mỗi class chỉ được chụp trên một mặt bàn hoặc một điều kiện ánh sáng.",
        "Kích thước bbox không phản ánh đúng dung tích khi khoảng cách camera thay đổi quá lớn.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "4.6. Đề xuất tiêu chí nghiệm thu", 2)
    add_para(doc, "Để nghiệm thu mô hình sau khi bổ sung dữ liệu thật, nên dùng tập kiểm thử độc lập với ít nhất 20% ảnh cho mỗi class. Tiêu chí khuyến nghị gồm: F1 score tổng thể đạt mức chấp nhận được theo yêu cầu môn học, confusion matrix không có class bị nhầm lẫn quá nhiều, centroid nằm gần vùng thân chai, và demo browser nhận diện ổn định trong nhiều điều kiện ánh sáng. Các số liệu cụ thể cần lấy trực tiếp từ Edge Impulse sau khi hoàn tất huấn luyện.")

    add_heading(doc, "4.7. Phân tích các kịch bản kiểm thử", 2)
    for text in [
        "Kịch bản kiểm thử đầu tiên là ảnh đơn chai trên nền đơn giản. Đây là kịch bản dễ nhất và dùng để xác nhận mô hình học đúng class cơ bản. Nếu mô hình còn nhầm mạnh trong kịch bản này, nguyên nhân thường là nhãn chưa nhất quán, class quá ít ảnh hoặc cấu hình impulse chưa phù hợp. Không nên chuyển sang kiểm thử nhiều chai khi mô hình chưa ổn định ở ảnh đơn chai.",
        "Kịch bản thứ hai là nhiều chai trong cùng một frame. Đây là nơi FOMO thể hiện ưu điểm vì có thể phát hiện nhiều centroid. Tuy nhiên, các chai không nên chồng lên nhau quá nhiều. Nếu hai chai sát nhau, heatmap có thể gộp thành một vùng activation hoặc centroid bị kéo vào giữa hai chai. Với băng chuyền thật, nên thiết kế cơ khí hoặc luồng đặt chai để khoảng cách giữa các chai đủ lớn.",
        "Kịch bản thứ ba là thay đổi ánh sáng và nền. Chai nhựa trong suốt dễ phản xạ ánh sáng, khiến nhãn dán hoặc đường viền thân chai khó nhìn. Nếu mô hình chỉ hoạt động ở ánh sáng phòng học nhưng lỗi dưới ánh sáng ngoài trời, cần bổ sung dữ liệu ngoài trời hoặc dùng augmentation brightness/contrast ở mức hợp lý. Kịch bản này rất quan trọng vì điểm thu gom rác thường có điều kiện ánh sáng không ổn định.",
        "Kịch bản thứ tư là kiểm thử realtime trên browser. Người kiểm thử nên đưa từng loại chai vào khung hình, ghi nhận nhãn dự đoán, score và vị trí centroid. Nếu score dao động mạnh theo từng frame, có thể cần tăng threshold hoặc dùng post-processing theo cửa sổ thời gian, chỉ xác nhận phát hiện khi cùng nhãn xuất hiện liên tiếp nhiều frame. Cách này giảm nhấp nháy kết quả khi demo.",
    ]:
        add_para(doc, text)

    add_heading(doc, "4.8. Rủi ro và biện pháp giảm thiểu", 2)
    for text in [
        "Rủi ro đầu tiên là dữ liệu chưa đại diện cho môi trường thật. Bộ ảnh chụp trong lớp học thường sạch và dễ hơn môi trường thu gom rác. Biện pháp giảm thiểu là tạo tập test riêng chụp ở môi trường gần thực tế, không dùng ảnh này trong training. Nếu model chỉ tốt trên training nhưng kém trên tập test thực tế, cần thu thêm dữ liệu thay vì chỉ tăng epoch.",
        "Rủi ro thứ hai là nhãn dung tích không luôn nhìn thấy. Một số chai bị xoay mặt không có nhãn về phía camera, hoặc nhãn bị rách/mờ. Trong trường hợp đó, hệ thống chỉ dựa vào hình dáng và kích thước tương đối, dễ nhầm hơn. Giải pháp là bố trí camera nhìn vào mặt nhãn khi có thể, hoặc kết hợp thêm cảm biến khoảng cách/cân nặng trong phiên bản sau để hỗ trợ phân loại.",
        "Rủi ro thứ ba là nhầm lẫn do phối cảnh. Cùng một chai 500ml ở gần camera có thể có bbox lớn hơn chai 1.5L ở xa camera. Vì vậy, metric bbox_area_ratio chỉ có ý nghĩa khi setup camera ổn định. Nếu hệ thống cho phép khoảng cách thay đổi nhiều, cần thêm bước chuẩn hóa bằng vùng quan sát cố định, marker tham chiếu hoặc dùng mô hình học trực tiếp từ nhiều khoảng cách khác nhau.",
        "Rủi ro cuối cùng là triển khai thiếu kiểm soát threshold. Threshold quá thấp làm tăng false positive, threshold quá cao làm bỏ sót chai. Sau khi có kết quả thật, nên khảo sát nhiều threshold và chọn giá trị cân bằng giữa precision và recall theo mục tiêu hệ thống. Nếu ưu tiên không bỏ sót chai, chọn threshold thấp hơn; nếu ưu tiên không kích hoạt sai cơ cấu phân loại, chọn threshold cao hơn.",
    ]:
        add_para(doc, text)


def chapter_5(doc: Document) -> None:
    add_heading(doc, "CHƯƠNG 5. TỔNG KẾT VÀ HƯỚNG PHÁT TRIỂN", 1)
    add_heading(doc, "5.1. Kết luận", 2)
    for text in [
        "Tiểu luận đã trình bày một thiết kế hệ thống phát hiện chai nhựa tái chế theo ba kích thước 500ml, 1L và 1.5L bằng mô hình FOMO trên Edge Impulse. Đề tài phù hợp với học phần Mạng cảm biến vì kết hợp node camera, xử lý tại biên, truyền kết quả rút gọn và dashboard/log phục vụ giám sát.",
        "Nội dung báo cáo làm rõ sự khác biệt giữa bounding box trong giai đoạn gán nhãn và centroid trong giai đoạn suy luận của FOMO. Đây là điểm quan trọng để tránh hiểu sai rằng FOMO có thể đo trực tiếp kích thước bounding box ở runtime. Với yêu cầu của đề tài, dung tích chai được xử lý như nhãn phân loại, còn bbox_area_ratio được dùng để phân tích chất lượng dữ liệu huấn luyện.",
        "Báo cáo cũng đưa ra pipeline triển khai rõ ràng: tạo project Edge Impulse của đề tài, upload dataset prototype, gán nhãn bounding box, thiết kế impulse, generate features, huấn luyện FOMO, đánh giá bằng metric thật, build deployment và chạy browser realtime. Phần cần bổ sung trong tương lai là kiểm thử với chai nhựa thật ngoài hiện trường, không phải chỉ runtime webcam.",
    ]:
        add_para(doc, text)

    add_heading(doc, "5.2. Hướng phát triển đề tài", 2)
    for item in [
        "Bổ sung dataset thật với nhiều thương hiệu, nhiều nền và nhiều điều kiện ánh sáng để tăng khả năng tổng quát.",
        "Tích hợp camera edge node như Raspberry Pi, ESP32-CAM hoặc thiết bị MCU có camera để kiểm thử ngoài trình duyệt.",
        "Thêm dashboard ghi log số lượng chai theo thời gian, xuất CSV và cảnh báo khi phát hiện sai nhãn nhiều lần.",
        "Nghiên cứu object tracking để tránh đếm trùng một chai qua nhiều frame liên tiếp.",
        "So sánh FOMO với MobileNetV2 SSD FPN hoặc YOLO-Pro nếu phiên bản sau cần bounding box đầy đủ ở runtime.",
        "Mở rộng class sang lon nhôm, chai thủy tinh, hộp giấy hoặc phân loại vật liệu tái chế theo nhiều nhóm hơn.",
    ]:
        add_bullet(doc, item)

    add_table_caption(doc, "Bảng 5.1. Tổng hợp kết quả đạt được và hướng phát triển")
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for i, h in enumerate(["Nội dung", "Trạng thái", "Ghi chú"]):
        table.cell(0, i).text = h
    rows = [
        ("Cấu trúc báo cáo theo mẫu PTIT", "Hoàn thành", "Đủ bìa, nhiệm vụ, chương 1-5, tài liệu tham khảo, phụ lục"),
        ("Thiết kế pipeline Edge Impulse/FOMO", "Hoàn thành", "Dựa trên project Edge Impulse của đề tài và tài liệu Edge Impulse"),
        ("Metric bbox_area_ratio", "Hoàn thành bằng dữ liệu prototype", "Mean: 500ml = 0.06330, 1L = 0.09584, 1.5L = 0.13828"),
        ("Training/confusion matrix", "Đã hoàn thành bằng Edge Impulse", "Validation F1 86.5%, test accuracy 86.67%, test F1 0.92"),
        ("Deployment/browser realtime", "Đã hoàn thành", "Build v12 C++ library và browser client đã chạy Inferencing"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val
    format_table(table, [5.0, 3.2, 7.8])


def references_and_appendix(doc: Document) -> None:
    add_heading(doc, "TÀI LIỆU THAM KHẢO", 1)
    refs = [
        "[1] Edge Impulse, \"FOMO,\" Edge Impulse Documentation, 2026. https://docs.edgeimpulse.com/studio/projects/learning-blocks/blocks/object-detection/fomo",
        "[2] Edge Impulse, \"Object detection with centroids,\" Edge Impulse Documentation, 2026. https://docs.edgeimpulse.com/tutorials/end-to-end/object-detection-centroids",
        "[3] Edge Impulse, \"Object detection,\" Edge Impulse Documentation, 2026. https://docs.edgeimpulse.com/studio/projects/learning-blocks/blocks/object-detection",
        "[4] Edge Impulse, \"Object detection image dataset annotation formats,\" Edge Impulse Documentation, 2026. https://docs.edgeimpulse.com/reference/data-ingestion/image-dataset-annotation-formats",
        "[5] Edge Impulse, \"Run WebAssembly library (browser),\" Edge Impulse Documentation, 2026. https://docs.edgeimpulse.com/docs/run-inference/webassembly/through-webassembly-browser",
        "[6] Edge Impulse, \"Deployment,\" Edge Impulse Documentation, 2026. https://docs.edgeimpulse.com/studio/projects/deployment",
        "[7] Wikimedia Commons, \"Logo PTIT University.png,\" 2021. https://commons.wikimedia.org/wiki/File:Logo_PTIT_University.png",
    ]
    for ref in refs:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.first_line_indent = Cm(-0.75)
        p.paragraph_format.line_spacing = 1.2
        p.paragraph_format.space_before = Pt(3)
        r = p.add_run(ref)
        set_run_font(r, 12)

    appendix = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(appendix)
    add_footer_page_number(appendix, fmt="decimal", start=1, prefix="PL-")
    add_heading(doc, "PHỤ LỤC", 1)
    add_heading(doc, "PHỤ LỤC 1. Cấu trúc thư mục repository đề xuất", 2)
    add_para(doc, "Repository GitHub/GitLab nên được chuẩn bị theo cấu trúc sau để giảng viên dễ kiểm tra. Link repo thật cần thay vào placeholder trong phần mở đầu trước khi nộp.", first_line=False)
    code_lines = [
        "N23DCCI034-PhamNgocQuangHuy-MangCamBien-FOMO/",
        "  README.md",
        "  report/",
        "    PhamNgocQuangHuy_final_cuoiky.docx",
        "  scripts/",
        "    build_final_tieuluan.py",
        "    make_synthetic_bottle_dataset.py",
        "    upload_synthetic_dataset_to_edge_impulse.py",
        "    edge_impulse_upload_bridge.py",
        "  data/",
        "    synthetic_bottle_dataset_edge_impulse.zip",
        "  screenshots/",
        "    08_real_uploaded_bottle_dataset.png",
        "    09_real_feature_generation.png",
        "    11_real_fomo_training_results.png",
        "    12_real_model_testing_results.png",
        "    14_real_deployment_build_success.png",
        "    15_real_browser_realtime_inferencing_full_redacted.png",
    ]
    for line in code_lines:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.left_indent = Cm(1.0)
        r = p.add_run(line)
        set_run_font(r, 11, name="Consolas")

    add_heading(doc, "PHỤ LỤC 2. Pseudocode hậu xử lý kết quả FOMO", 2)
    pseudocode = [
        "Input: frame từ camera",
        "Resize frame theo cấu hình impulse",
        "result = run_classifier(frame)",
        "for detection in result.bounding_boxes hoặc result centroids:",
        "    if detection.score >= threshold:",
        "        ghi log {label, x, y, score, timestamp}",
        "        tăng bộ đếm theo label nếu chưa đếm đối tượng này trong cửa sổ thời gian ngắn",
        "hiển thị label và centroid lên dashboard hoặc browser demo",
    ]
    for line in pseudocode:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.left_indent = Cm(0.7)
        r = p.add_run(line)
        set_run_font(r, 11.5, name="Consolas")

    add_heading(doc, "PHỤ LỤC 3. Danh sách ảnh minh chứng đã có và cần bổ sung", 2)
    for item in [
        "Đã có ảnh Data acquisition thể hiện 120 mẫu, 90 training và 30 testing.",
        "Đã có ảnh feature generation cho 90 mẫu training và ba class 500ml, 1L, 1.5L.",
        "Đã có ảnh kết quả training, confusion matrix và model testing.",
        "Đã có ảnh deployment build thành công và browser realtime client đang Inferencing.",
        "Nên bổ sung thêm ảnh demo với chai thật trước camera nếu giảng viên yêu cầu nghiệm thu ngoài hiện trường.",
    ]:
        add_bullet(doc, item)


def add_update_fields_setting(doc: Document) -> None:
    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def main() -> None:
    doc = Document()
    configure_styles(doc)
    configure_section(doc.sections[0])
    clear_footer(doc.sections[0])
    add_update_fields_setting(doc)

    cover_page(doc, secondary=False)
    front = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(front)
    add_footer_page_number(front, fmt="lowerRoman", start=1)
    cover_page(doc, secondary=True)
    add_page_break(doc)
    task_page(doc)
    add_page_break(doc)
    declaration_pages(doc)
    front_matter_lists(doc)

    main_sec = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(main_sec)
    add_footer_page_number(main_sec, fmt="decimal", start=1)
    chapter_1(doc)
    add_page_break(doc)
    chapter_2(doc)
    add_page_break(doc)
    chapter_3(doc)
    add_page_break(doc)
    chapter_4(doc)
    add_page_break(doc)
    chapter_5(doc)
    add_page_break(doc)
    references_and_appendix(doc)

    doc.core_properties.author = "Pham Ngoc Quang Huy"
    doc.core_properties.title = "Phat hien chai nhua tai che theo size 500ml 1L 1.5L"
    doc.core_properties.subject = "Mang Cam Bien"
    doc.save(OUT_DOCX)
    print(OUT_DOCX)


if __name__ == "__main__":
    main()
